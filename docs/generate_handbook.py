from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
OUT='docs/Edward_Wonder_Mordey_Portfolio_Owners_Handbook.docx'
INK='1D2924'; GREEN='335A46'; MUTED='637168'; PALE='F2F5F2'; BORDER='D6DED8'
def setfont(r,size=9,color=INK,bold=None):
 r.font.name='Aptos';r.font.size=Pt(size);r.font.color.rgb=RGBColor.from_string(color)
 if bold is not None:r.bold=bold
def shade(c,col):
 p=c._tc.get_or_add_tcPr();s=OxmlElement('w:shd');s.set(qn('w:fill'),col);p.append(s)
def border(c):
 p=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
 for edge in ('top','left','bottom','right','insideH','insideV'):
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),BORDER);b.append(e)
 p.append(b)
def H(s,l=1):
 p=d.add_heading(s,l);p.paragraph_format.keep_with_next=True;return p
def P(s):
 p=d.add_paragraph(s);p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.12
 for r in p.runs:setfont(r)
 return p
def B(xs):
 for s in xs:
  p=d.add_paragraph(style='List Bullet');p.paragraph_format.space_after=Pt(3);p.add_run(s)
def S(xs):
 for i,s in enumerate(xs,1):
  p=d.add_paragraph();p.paragraph_format.left_indent=Cm(.55);p.paragraph_format.first_line_indent=Cm(-.55);p.paragraph_format.space_after=Pt(4);p.add_run(str(i)+'. '+s)
def T(head,rows):
 t=d.add_table(rows=1,cols=len(head));t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for j,x in enumerate(head):
  c=t.rows[0].cells[j];c.text=x;shade(c,'263C32');border(c);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for r in c.paragraphs[0].runs:setfont(r,8,'FFFFFF',True)
 for i,row in enumerate(rows):
  cs=t.add_row().cells
  for j,x in enumerate(row):
   cs[j].text=str(x);shade(cs[j],'FFFFFF' if i%2==0 else PALE);border(cs[j]);cs[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in cs[j].paragraphs:
    for r in p.runs:setfont(r,8)
 d.add_paragraph()
d=Document();sec=d.sections[0];sec.page_width=Cm(21.59);sec.page_height=Cm(27.94);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8);sec.different_first_page_header_footer=True
st=d.styles;st['Normal'].font.name='Aptos';st['Normal'].font.size=Pt(9);st['Normal'].font.color.rgb=RGBColor.from_string(INK)
for n,z in [('Title',29),('Heading 1',18),('Heading 2',12),('Heading 3',10)]:
 st[n].font.name='Aptos Display';st[n].font.size=Pt(z);st[n].font.bold=True;st[n].font.color.rgb=RGBColor.from_string('000000')
 if n=='Heading 1':st[n].paragraph_format.page_break_before=True
title_ppr=st['Title']._element.get_or_add_pPr()
for node in list(title_ppr):
 if node.tag.endswith('pBdr'):title_ppr.remove(node)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=f.add_run('Edward Wonder Mordey  |  Portfolio Website Owner’s Handbook');setfont(r,8,MUTED)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(80);r=p.add_run('EDWARD WONDER MORDEY');setfont(r,14,GREEN,True)
p=d.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('Portfolio Website Owner’s Handbook')
for r in p.runs:setfont(r,29,'000000',True)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('IT · Engineering · Service\nBuilding systems with purpose.');setfont(r,14,GREEN)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(30);r=p.add_run('A practical guide to updating and maintaining the professional portfolio.');setfont(r,11,MUTED)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(70);r=p.add_run('https://edward-wonder-mordey.github.io/portfolio/');setfont(r,10,GREEN)
d.add_page_break()
H('1. Purpose and safe editing')
P('This handbook explains how the portfolio is organized, where to change common content, and how to publish updates. The website is a public record of technical work, education, communication and service. Keep every public statement accurate and remove confidential details from examples.')
H('Before each update',2);S(['Open the live site and identify the exact section, card, image or link to change.','Edit the main branch in the portfolio repository. Preserve the current structure and accessibility text.','Review the change, all links and images, then run a production build when working locally.','Wait for GitHub Actions to complete successfully. Check the live page on desktop and mobile.'])
T(['Change level','Examples','Review needed'],[('Content','Text, contact details, project description or image','Verify every fact, date and link.'),('Design','Section order, spacing, colors or mobile behavior','Test keyboard use, reduced motion and phone layout.'),('Technical','Dependencies, workflow, build settings or security','Back up and review the full change; run build and deployment checks.')])
H('2. Quick file finder')
T(['Change','File or folder','How it is used'],[
('Main content and page sections','src/App.tsx','Main React and TypeScript page: hero, About, Work, Journey, Academics, Faith & Growth, Contact.'),
('Main visual system','src/index.css','Shared layout, colors, typography, navigation and responsive styles.'),
('Academic styling','src/academic.css','Academic profile and credential presentation.'),
('Faith and Growth styling','src/editorial-growth.css','Reflective and professional growth section.'),
('Detailed case studies','public/case-studies.html','Long-form descriptions linked from project cards.'),
('Project imagery','public/images/','Image assets; update the matching path in src/App.tsx.'),
('Certificates','public/certificates/ and public/','Public downloadable evidence. Publish only documents approved for public viewing.'),
('SQL, presentations, documents, reflections','public/sql/, public/presentations/, public/documents/, public/reflections/','Downloadable samples linked from sections or case studies.'),
('Browser metadata','index.html','Base page title, description and social preview metadata.'),
('Build and hosting','package.json, vite.config.ts, .github/workflows/deploy.yml','Technical setup; routine text edits usually do not need changes here.')
])
H('3. How publishing works')
P('This repository uses React, TypeScript and Vite. The GitHub Actions workflow builds the site and publishes the generated dist directory to GitHub Pages.')
S(['A visitor opens the published portfolio URL.','index.html loads the built React application.','src/App.tsx renders the page and references public assets.','A change pushed to main starts .github/workflows/deploy.yml.','After a successful deployment, GitHub Pages serves the new build.'])
T(['Item','Current value'],[('Repository','github.com/Edward-Wonder-Mordey/portfolio'),('Live site','https://edward-wonder-mordey.github.io/portfolio/'),('Source branch','main'),('Local development','npm run dev'),('Production build','npm run build')])
H('4. Sections and content')
T(['Section','Purpose','Update guidance'],[
('Hero','Name and professional positioning.','Preserve “IT · Engineering · Service” and the headline unless the personal brand changes.'),
('About me','Technical focus and working approach.','Describe current strengths plainly; avoid unsupported claims.'),
('Selected Work','Project cards, tags and case study links.','Use genuine work and distinguish coursework, paid work, personal projects and service.'),
('Journey','Experience and service timeline.','Keep dates, role titles and organization names correct.'),
('Academic profile','Programs, records, credentials and learning.','Check source records; label active study accurately.'),
('Faith & Growth','Personal reflection and professional development.','Review what personal details should remain public.'),
('Contact','Email and public profiles.','Test each link and update the README if public details change.')
])
P('Most visible copy and project entries are in src/App.tsx. Search for the exact title or section id before editing so the correct component is changed.')
H('5. Add or revise a project')
S(['In src/App.tsx, find the Selected Work section or the project title. Follow the structure of a neighboring entry.','State the context, your role, the work performed and the real output. Identify coursework or volunteer work as such.','Place an image in the appropriate public/images folder. Use a clear filename and update the referenced path.','If the card links to a detailed case study, update public/case-studies.html and keep its anchor/link consistent.','Write useful alt text, check the layout on mobile, and test the link after deployment.'])
T(['Current portfolio work themes','Evidence focus'],[
('PromisedLand Data Center','Hybrid cloud architecture and infrastructure proposal.'),
('Eragonova Enterprise','Digital enterprise, systems, website and service work.'),
('Linux Deployment Proposal','Linux, storage and security planning.'),
('Data Analytics & Reporting','Power BI, Tableau, SQL, PostgreSQL and Excel.'),
('Networking & Security Labs','Packet Tracer, Nmap, Wireshark and Linux exercises.'),
('Stake Welfare & Self-Reliance Plan · 2026','Planning, coordination and reporting.'),
('Graphic Design & Visual Identity','Branding, layout and visual communication.'),
('Data Detective Research Presentation','Research analysis and presentation design.'),
('Kumasi YSA Quick Start Guide','Technical writing and onboarding.'),
('Communication Development Plan','Listening, reflection and professional growth.')
])
H('6. Academic records and public evidence')
P('Academic qualifications, dates, grades and credential names are sensitive to error. Check the underlying record before publishing or changing them.')
B(['Add a downloadable credential to public/certificates/ and update its exact link in src/App.tsx.','The repository also holds selected documents at public/ and subject examples under public/images/data-analytics/ and public/images/presentations/.','Do not publish student IDs, addresses, account details or private transcript information.','Show courses in progress as in progress, not completed credentials.','Test every download path after changing a filename.'])
P('Anything under public/ can be accessed directly by visitors. Do not store private drafts, client material, passwords, API keys or confidential records there.')
H('7. Brand and writing')
T(['Element','Current standard'],[('Name','Edward Wonder Mordey.'),('Positioning','IT · Engineering · Service.'),('Headline','Building systems with purpose.'),('Voice','Clear, grounded, accurate and service-oriented.'),('Approach','Explain what work demonstrates without overstating results.')])
H('8. Privacy and contact')
B(['Publish only information you intend to make public.','Remove client names, private screenshots, signatures and identifying data unless you have permission.','Never commit passwords, access tokens, API keys or private configuration.','Keep the website contact details and README consistent when a public link changes.'])
H('9. Publish with GitHub')
H('Small change from GitHub',2);S(['Open the portfolio repository and navigate to the exact file.','Use the editor for a small text correction, or upload an asset into the correct public folder.','Review the diff and commit the change to main.','Open Actions and wait for the deployment workflow to succeed.','Visit the live URL and test the changed section, downloads, links and mobile layout.'])
H('Local development',2);P('Use Node.js 20 or newer and npm. In a local clone, run npm install, then npm run dev. Before deployment, run npm run build and review the generated site with npm run preview. Do not edit dist as the source of truth.')
H('10. Troubleshooting')
T(['Problem','Likely cause','What to check'],[
('Update is not visible','Deployment has not finished or the wrong file was edited.','Check Actions on main, then refresh the public URL.'),
('Broken image/download','Path or filename does not match exactly.','Compare capitalization, spaces and punctuation in the public path.'),
('Build failed','Invalid JSX/TypeScript or import.','Review the Actions log or run npm run build.'),
('Layout breaks on phone','A style change was only checked on desktop.','Review responsive rules in src/index.css and test at phone width.'),
('Case study link fails','Card URL and case-study anchor differ.','Check the href and matching section in public/case-studies.html.'),
('Private file is public','File was added under public/.','Remove it and assess whether repository history also needs cleanup.')])
H('11. Routine maintenance')
B(['Check the live navigation and contact links after significant edits.','Keep experience dates, academic details and credential labels current.','Check file access and permissions for every new piece of evidence.','Check GitHub Actions after publishing.','Update this handbook and README when site structure or deployment changes.','Remove outdated or misleading material.'])
H('12. Repository reference')
T(['Purpose','Location'],[('Main page','src/App.tsx'),('Case studies','public/case-studies.html'),('Styles','src/index.css, src/academic.css, src/editorial-growth.css'),('Deployment','GitHub Actions: .github/workflows/deploy.yml'),('Source branch','main'),('Live address','https://edward-wonder-mordey.github.io/portfolio/')])
P('Keep a short change record in GitHub commit messages. When a major section, build step or deployment method changes, update this handbook and its generator.')
for _tbl in d.tables:
 for _ri,_row in enumerate(_tbl.rows):
  _trpr=_row._tr.get_or_add_trPr();_trpr.append(OxmlElement('w:cantSplit'))
  if _ri==0:
   _hdr=OxmlElement('w:tblHeader');_hdr.set(qn('w:val'),'true');_trpr.append(_hdr)
props=d.core_properties;props.title='Portfolio Website Owners Handbook';props.subject='Guide to updating and maintaining Edward Wonder Mordey’s portfolio website';props.author='Edward Wonder Mordey'
Path('docs').mkdir(parents=True,exist_ok=True);d.save(OUT);print(OUT)
